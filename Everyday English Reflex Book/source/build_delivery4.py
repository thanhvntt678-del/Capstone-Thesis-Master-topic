# -*- coding: utf-8 -*-
"""Builds Delivery #4 (Lessons 0024 onward) — same pattern as build_delivery3.py:
no pre-approved base docx, cross-checks duplicates against the WHOLE book so far
(Deliveries #1, #2, #3 plus whatever's in this delivery)."""
import sys
sys.path.insert(0, '.')
import docx
from docx.shared import Pt, RGBColor
from lesson_builder import (
    add_title_block, add_lesson_header_table, add_blank_spacer,
    add_intro_paragraph, add_dialogue_line, qc_report, lesson_word_count,
    new_section_setup, add_page_number_field,
)
from build_combined import load_lesson0001_as_dict, cross_lesson_duplicate_check
from lesson0002 import LESSON_0002
from lesson0003 import LESSON_0003
from lesson0004 import LESSON_0004
from lesson0005 import LESSON_0005
from lesson0006 import LESSON_0006
from lesson0007 import LESSON_0007
from lesson0008 import LESSON_0008
from lesson0009 import LESSON_0009
from lesson0010 import LESSON_0010
from lesson0011 import LESSON_0011
from lesson0012 import LESSON_0012
from lesson0013 import LESSON_0013
from lesson0014 import LESSON_0014
from lesson0015 import LESSON_0015
from lesson0016 import LESSON_0016
from lesson0017 import LESSON_0017
from lesson0018 import LESSON_0018
from lesson0019 import LESSON_0019
from lesson0020 import LESSON_0020
from lesson0021 import LESSON_0021
from lesson0022 import LESSON_0022
from lesson0023 import LESSON_0023
from lesson0024 import LESSON_0024
from lesson0025 import LESSON_0025
from lesson0026 import LESSON_0026
from lesson0027 import LESSON_0027
from lesson0028 import LESSON_0028
from lesson0029 import LESSON_0029
from lesson0030 import LESSON_0030

DELIVERY1 = [LESSON_0002, LESSON_0003, LESSON_0004, LESSON_0005, LESSON_0006, LESSON_0007]
DELIVERY2 = [LESSON_0008, LESSON_0009, LESSON_0010, LESSON_0011, LESSON_0012, LESSON_0013, LESSON_0014, LESSON_0015]
DELIVERY3 = [LESSON_0016, LESSON_0017, LESSON_0018, LESSON_0019, LESSON_0020, LESSON_0021, LESSON_0022, LESSON_0023]

LESSONS_DELIVERY4 = [LESSON_0024, LESSON_0025, LESSON_0026, LESSON_0027, LESSON_0028, LESSON_0029, LESSON_0030]  # append 0031, ... here as they're written

def main():
    doc = docx.Document()
    new_section_setup(doc)

    for idx, lesson in enumerate(LESSONS_DELIVERY4):
        if idx != 0:
            doc.add_page_break()
        add_title_block(doc)
        add_lesson_header_table(doc, lesson['cefr'], lesson['lesson_id'], lesson['en_title'], lesson['vi_title'])
        add_blank_spacer(doc)
        add_intro_paragraph(doc, lesson['intro_en'], lesson['intro_vi'])
        for speaker, en, vi in lesson['turns']:
            add_dialogue_line(doc, speaker, en, vi)

    lesson_range = f"{LESSONS_DELIVERY4[0]['lesson_id']}-{LESSONS_DELIVERY4[-1]['lesson_id']}"

    sec = doc.sections[0]
    ftr = sec.footer
    p = ftr.paragraphs[0]
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"EVERYDAY ENGLISH REFLEX  •  Lessons {lesson_range}  •  page ")
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    r.font.size = Pt(8)
    add_page_number_field(p)

    out_path = f"EVERYDAY_ENGLISH_REFLEX_LESSONS_{lesson_range}.docx"
    doc.save(out_path)
    print("Saved:", out_path)

    print("\n=== PER-LESSON QC (Delivery #4) ===")
    for lesson in LESSONS_DELIVERY4:
        r = qc_report(lesson)
        print(f"Lesson {r['lesson_id']}: English words={r['english_words']}, turns={r['turns']}, "
              f"speakers={sorted(r['speakers'])}, duplicate_lines={r['duplicate_lines']}")

    lesson0001 = load_lesson0001_as_dict("lesson0001_approved.docx")
    WHOLE_BOOK = [lesson0001] + DELIVERY1 + DELIVERY2 + DELIVERY3 + LESSONS_DELIVERY4
    dups = cross_lesson_duplicate_check(WHOLE_BOOK)
    print(f"\nCross-lesson duplicate English lines (whole book, 0001 through {LESSONS_DELIVERY4[-1]['lesson_id']}):", dups)

    total_words = sum(lesson_word_count(l) for l in LESSONS_DELIVERY4)
    total_turns = sum(len(l['turns']) for l in LESSONS_DELIVERY4) + len(LESSONS_DELIVERY4)
    est_pages = total_turns / (133 / 5)
    print(f"\nTotal English learning words, Delivery #4 ({lesson_range}):", total_words)
    print(f"Structural page estimate for Delivery #4 so far: {est_pages:.1f} pages")

if __name__ == "__main__":
    main()
