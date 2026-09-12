# -*- coding: utf-8 -*-
"""Builds Delivery #19 (Lessons 0154 onward) — same pattern as build_delivery18.py:
no pre-approved base docx, cross-checks duplicates against the WHOLE book so far
(Deliveries #1-#18 plus whatever's in this delivery)."""
import sys
sys.path.insert(0, '.')
import docx
from docx.shared import Pt, RGBColor
from lesson_builder import (
    add_title_block, add_lesson_header_table, add_blank_spacer,
    add_intro_paragraph, add_dialogue_line, qc_report, lesson_word_count,
    new_section_setup, add_page_number_field,
)
from build_combined import load_lesson0001_as_dict, cross_lesson_duplicate_check
from lesson0002 import LESSON_0002
from lesson0003 import LESSON_0003
from lesson0004 import LESSON_0004
from lesson0005 import LESSON_0005
from lesson0006 import LESSON_0006
from lesson0007 import LESSON_0007
from lesson0008 import LESSON_0008
from lesson0009 import LESSON_0009
from lesson0010 import LESSON_0010
from lesson0011 import LESSON_0011
from lesson0012 import LESSON_0012
from lesson0013 import LESSON_0013
from lesson0014 import LESSON_0014
from lesson0015 import LESSON_0015
from lesson0016 import LESSON_0016
from lesson0017 import LESSON_0017
from lesson0018 import LESSON_0018
from lesson0019 import LESSON_0019
from lesson0020 import LESSON_0020
from lesson0021 import LESSON_0021
from lesson0022 import LESSON_0022
from lesson0023 import LESSON_0023
from lesson0024 import LESSON_0024
from lesson0025 import LESSON_0025
from lesson0026 import LESSON_0026
from lesson0027 import LESSON_0027
from lesson0028 import LESSON_0028
from lesson0029 import LESSON_0029
from lesson0030 import LESSON_0030
from lesson0031 import LESSON_0031
from lesson0032 import LESSON_0032
from lesson0033 import LESSON_0033
from lesson0034 import LESSON_0034
from lesson0035 import LESSON_0035
from lesson0036 import LESSON_0036
from lesson0037 import LESSON_0037
from lesson0038 import LESSON_0038
from lesson0039 import LESSON_0039
from lesson0040 import LESSON_0040
from lesson0041 import LESSON_0041
from lesson0042 import LESSON_0042
from lesson0043 import LESSON_0043
from lesson0044 import LESSON_0044
from lesson0045 import LESSON_0045
from lesson0046 import LESSON_0046
from lesson0047 import LESSON_0047
from lesson0048 import LESSON_0048
from lesson0049 import LESSON_0049
from lesson0050 import LESSON_0050
from lesson0051 import LESSON_0051
from lesson0052 import LESSON_0052
from lesson0053 import LESSON_0053
from lesson0054 import LESSON_0054
from lesson0055 import LESSON_0055
from lesson0056 import LESSON_0056
from lesson0057 import LESSON_0057
from lesson0058 import LESSON_0058
from lesson0059 import LESSON_0059
from lesson0060 import LESSON_0060
from lesson0061 import LESSON_0061
from lesson0062 import LESSON_0062
from lesson0063 import LESSON_0063
from lesson0064 import LESSON_0064
from lesson0065 import LESSON_0065
from lesson0066 import LESSON_0066
from lesson0067 import LESSON_0067
from lesson0068 import LESSON_0068
from lesson0069 import LESSON_0069
from lesson0070 import LESSON_0070
from lesson0071 import LESSON_0071
from lesson0072 import LESSON_0072
from lesson0073 import LESSON_0073
from lesson0074 import LESSON_0074
from lesson0075 import LESSON_0075
from lesson0076 import LESSON_0076
from lesson0077 import LESSON_0077
from lesson0078 import LESSON_0078
from lesson0079 import LESSON_0079
from lesson0080 import LESSON_0080
from lesson0081 import LESSON_0081
from lesson0082 import LESSON_0082
from lesson0083 import LESSON_0083
from lesson0084 import LESSON_0084
from lesson0085 import LESSON_0085
from lesson0086 import LESSON_0086
from lesson0087 import LESSON_0087
from lesson0088 import LESSON_0088
from lesson0089 import LESSON_0089
from lesson0090 import LESSON_0090
from lesson0091 import LESSON_0091
from lesson0092 import LESSON_0092
from lesson0093 import LESSON_0093
from lesson0094 import LESSON_0094
from lesson0095 import LESSON_0095
from lesson0096 import LESSON_0096
from lesson0097 import LESSON_0097
from lesson0098 import LESSON_0098
from lesson0099 import LESSON_0099
from lesson0100 import LESSON_0100
from lesson0101 import LESSON_0101
from lesson0102 import LESSON_0102
from lesson0103 import LESSON_0103
from lesson0104 import LESSON_0104
from lesson0105 import LESSON_0105
from lesson0106 import LESSON_0106
from lesson0107 import LESSON_0107
from lesson0108 import LESSON_0108
from lesson0109 import LESSON_0109
from lesson0110 import LESSON_0110
from lesson0111 import LESSON_0111
from lesson0112 import LESSON_0112
from lesson0113 import LESSON_0113
from lesson0114 import LESSON_0114
from lesson0115 import LESSON_0115
from lesson0116 import LESSON_0116
from lesson0117 import LESSON_0117
from lesson0118 import LESSON_0118
from lesson0119 import LESSON_0119
from lesson0120 import LESSON_0120
from lesson0121 import LESSON_0121
from lesson0122 import LESSON_0122
from lesson0123 import LESSON_0123
from lesson0124 import LESSON_0124
from lesson0125 import LESSON_0125
from lesson0126 import LESSON_0126
from lesson0127 import LESSON_0127
from lesson0128 import LESSON_0128
from lesson0129 import LESSON_0129
from lesson0130 import LESSON_0130
from lesson0131 import LESSON_0131
from lesson0132 import LESSON_0132
from lesson0133 import LESSON_0133
from lesson0134 import LESSON_0134
from lesson0135 import LESSON_0135
from lesson0136 import LESSON_0136
from lesson0137 import LESSON_0137
from lesson0138 import LESSON_0138
from lesson0139 import LESSON_0139
from lesson0140 import LESSON_0140
from lesson0141 import LESSON_0141
from lesson0142 import LESSON_0142
from lesson0143 import LESSON_0143
from lesson0144 import LESSON_0144
from lesson0145 import LESSON_0145
from lesson0146 import LESSON_0146
from lesson0147 import LESSON_0147
from lesson0148 import LESSON_0148
from lesson0149 import LESSON_0149
from lesson0150 import LESSON_0150
from lesson0151 import LESSON_0151
from lesson0152 import LESSON_0152
from lesson0153 import LESSON_0153
from lesson0154 import LESSON_0154
from lesson0155 import LESSON_0155
from lesson0156 import LESSON_0156
from lesson0157 import LESSON_0157
from lesson0158 import LESSON_0158
from lesson0159 import LESSON_0159
from lesson0160 import LESSON_0160
from lesson0161 import LESSON_0161
from lesson0162 import LESSON_0162

DELIVERY1 = [LESSON_0002, LESSON_0003, LESSON_0004, LESSON_0005, LESSON_0006, LESSON_0007]
DELIVERY2 = [LESSON_0008, LESSON_0009, LESSON_0010, LESSON_0011, LESSON_0012, LESSON_0013, LESSON_0014, LESSON_0015]
DELIVERY3 = [LESSON_0016, LESSON_0017, LESSON_0018, LESSON_0019, LESSON_0020, LESSON_0021, LESSON_0022, LESSON_0023]
DELIVERY4 = [LESSON_0024, LESSON_0025, LESSON_0026, LESSON_0027, LESSON_0028, LESSON_0029, LESSON_0030, LESSON_0031]
DELIVERY5 = [LESSON_0032, LESSON_0033, LESSON_0034, LESSON_0035, LESSON_0036, LESSON_0037, LESSON_0038, LESSON_0039]
DELIVERY6 = [LESSON_0040, LESSON_0041, LESSON_0042, LESSON_0043, LESSON_0044, LESSON_0045, LESSON_0046, LESSON_0047]
DELIVERY7 = [LESSON_0048, LESSON_0049, LESSON_0050, LESSON_0051, LESSON_0052, LESSON_0053, LESSON_0054, LESSON_0055]
DELIVERY8 = [LESSON_0056, LESSON_0057, LESSON_0058, LESSON_0059, LESSON_0060, LESSON_0061, LESSON_0062, LESSON_0063]
DELIVERY9 = [LESSON_0064, LESSON_0065, LESSON_0066, LESSON_0067, LESSON_0068, LESSON_0069, LESSON_0070, LESSON_0071, LESSON_0072]
DELIVERY10 = [LESSON_0073, LESSON_0074, LESSON_0075, LESSON_0076, LESSON_0077, LESSON_0078, LESSON_0079, LESSON_0080, LESSON_0081]
DELIVERY11 = [LESSON_0082, LESSON_0083, LESSON_0084, LESSON_0085, LESSON_0086, LESSON_0087, LESSON_0088, LESSON_0089, LESSON_0090]
DELIVERY12 = [LESSON_0091, LESSON_0092, LESSON_0093, LESSON_0094, LESSON_0095, LESSON_0096, LESSON_0097, LESSON_0098, LESSON_0099]
DELIVERY13 = [LESSON_0100, LESSON_0101, LESSON_0102, LESSON_0103, LESSON_0104, LESSON_0105, LESSON_0106, LESSON_0107, LESSON_0108]
DELIVERY14 = [LESSON_0109, LESSON_0110, LESSON_0111, LESSON_0112, LESSON_0113, LESSON_0114, LESSON_0115, LESSON_0116, LESSON_0117]
DELIVERY15 = [LESSON_0118, LESSON_0119, LESSON_0120, LESSON_0121, LESSON_0122, LESSON_0123, LESSON_0124, LESSON_0125, LESSON_0126]
DELIVERY16 = [LESSON_0127, LESSON_0128, LESSON_0129, LESSON_0130, LESSON_0131, LESSON_0132, LESSON_0133, LESSON_0134, LESSON_0135]
DELIVERY17 = [LESSON_0136, LESSON_0137, LESSON_0138, LESSON_0139, LESSON_0140, LESSON_0141, LESSON_0142, LESSON_0143, LESSON_0144]
DELIVERY18 = [LESSON_0145, LESSON_0146, LESSON_0147, LESSON_0148, LESSON_0149, LESSON_0150, LESSON_0151, LESSON_0152, LESSON_0153]

LESSONS_DELIVERY19 = [LESSON_0154, LESSON_0155, LESSON_0156, LESSON_0157, LESSON_0158, LESSON_0159, LESSON_0160, LESSON_0161, LESSON_0162]  # 9 lessons closes near ~37 pages; 0163-0169 open Delivery #20

def main():
    doc = docx.Document()
    new_section_setup(doc)

    for idx, lesson in enumerate(LESSONS_DELIVERY19):
        if idx != 0:
            doc.add_page_break()
        add_title_block(doc)
        add_lesson_header_table(doc, lesson['cefr'], lesson['lesson_id'], lesson['en_title'], lesson['vi_title'])
        add_blank_spacer(doc)
        add_intro_paragraph(doc, lesson['intro_en'], lesson['intro_vi'])
        for speaker, en, vi in lesson['turns']:
            add_dialogue_line(doc, speaker, en, vi)

    lesson_range = f"{LESSONS_DELIVERY19[0]['lesson_id']}-{LESSONS_DELIVERY19[-1]['lesson_id']}"

    sec = doc.sections[0]
    ftr = sec.footer
    p = ftr.paragraphs[0]
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"EVERYDAY ENGLISH REFLEX  •  Lessons {lesson_range}  •  page ")
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    r.font.size = Pt(8)
    add_page_number_field(p)

    out_path = f"EVERYDAY_ENGLISH_REFLEX_LESSONS_{lesson_range}.docx"
    doc.save(out_path)
    print("Saved:", out_path)

    print("\n=== PER-LESSON QC (Delivery #19) ===")
    for lesson in LESSONS_DELIVERY19:
        r = qc_report(lesson)
        print(f"Lesson {r['lesson_id']}: English words={r['english_words']}, turns={r['turns']}, "
              f"speakers={sorted(r['speakers'])}, duplicate_lines={r['duplicate_lines']}")

    lesson0001 = load_lesson0001_as_dict("lesson0001_approved.docx")
    WHOLE_BOOK = [lesson0001] + DELIVERY1 + DELIVERY2 + DELIVERY3 + DELIVERY4 + DELIVERY5 + DELIVERY6 + DELIVERY7 + DELIVERY8 + DELIVERY9 + DELIVERY10 + DELIVERY11 + DELIVERY12 + DELIVERY13 + DELIVERY14 + DELIVERY15 + DELIVERY16 + DELIVERY17 + DELIVERY18 + LESSONS_DELIVERY19
    dups = cross_lesson_duplicate_check(WHOLE_BOOK)
    print(f"\nCross-lesson duplicate English lines (whole book, 0001 through {LESSONS_DELIVERY19[-1]['lesson_id']}):", dups)

    total_words = sum(lesson_word_count(l) for l in LESSONS_DELIVERY19)
    total_turns = sum(len(l['turns']) for l in LESSONS_DELIVERY19) + len(LESSONS_DELIVERY19)
    est_pages = total_turns / (133 / 5)
    print(f"\nTotal English learning words, Delivery #19 ({lesson_range}):", total_words)
    print(f"Structural page estimate for Delivery #19 so far: {est_pages:.1f} pages")

if __name__ == "__main__":
    main()
