# -*- coding: utf-8 -*-
"""Builds Delivery #13 (Lessons 0100 onward) — same pattern as build_delivery12.py:
no pre-approved base docx, cross-checks duplicates against the WHOLE book so far
(Deliveries #1-#12 plus whatever's in this delivery)."""
import sys
sys.path.insert(0, '.')
import docx
from docx.shared import Pt, RGBColor
from lesson_builder import (
    add_title_block, add_lesson_header_table, add_blank_spacer,
    add_intro_paragraph, add_dialogue_line, qc_report, lesson_word_count,
    new_section_setup, add_page_number_field,
)
from build_combined import load_lesson0001_as_dict, cross_lesson_duplicate_check
from lesson0002 import LESSON_0002
from lesson0003 import LESSON_0003
from lesson0004 import LESSON_0004
from lesson0005 import LESSON_0005
from lesson0006 import LESSON_0006
from lesson0007 import LESSON_0007
from lesson0008 import LESSON_0008
from lesson0009 import LESSON_0009
from lesson0010 import LESSON_0010
from lesson0011 import LESSON_0011
from lesson0012 import LESSON_0012
from lesson0013 import LESSON_0013
from lesson0014 import LESSON_0014
from lesson0015 import LESSON_0015
from lesson0016 import LESSON_0016
from lesson0017 import LESSON_0017
from lesson0018 import LESSON_0018
from lesson0019 import LESSON_0019
from lesson0020 import LESSON_0020
from lesson0021 import LESSON_0021
from lesson0022 import LESSON_0022
from lesson0023 import LESSON_0023
from lesson0024 import LESSON_0024
from lesson0025 import LESSON_0025
from lesson0026 import LESSON_0026
from lesson0027 import LESSON_0027
from lesson0028 import LESSON_0028
from lesson0029 import LESSON_0029
from lesson0030 import LESSON_0030
from lesson0031 import LESSON_0031
from lesson0032 import LESSON_0032
from lesson0033 import LESSON_0033
from lesson0034 import LESSON_0034
from lesson0035 import LESSON_0035
from lesson0036 import LESSON_0036
from lesson0037 import LESSON_0037
from lesson0038 import LESSON_0038
from lesson0039 import LESSON_0039
from lesson0040 import LESSON_0040
from lesson0041 import LESSON_0041
from lesson0042 import LESSON_0042
from lesson0043 import LESSON_0043
from lesson0044 import LESSON_0044
from lesson0045 import LESSON_0045
from lesson0046 import LESSON_0046
from lesson0047 import LESSON_0047
from lesson0048 import LESSON_0048
from lesson0049 import LESSON_0049
from lesson0050 import LESSON_0050
from lesson0051 import LESSON_0051
from lesson0052 import LESSON_0052
from lesson0053 import LESSON_0053
from lesson0054 import LESSON_0054
from lesson0055 import LESSON_0055
from lesson0056 import LESSON_0056
from lesson0057 import LESSON_0057
from lesson0058 import LESSON_0058
from lesson0059 import LESSON_0059
from lesson0060 import LESSON_0060
from lesson0061 import LESSON_0061
from lesson0062 import LESSON_0062
from lesson0063 import LESSON_0063
from lesson0064 import LESSON_0064
from lesson0065 import LESSON_0065
from lesson0066 import LESSON_0066
from lesson0067 import LESSON_0067
from lesson0068 import LESSON_0068
from lesson0069 import LESSON_0069
from lesson0070 import LESSON_0070
from lesson0071 import LESSON_0071
from lesson0072 import LESSON_0072
from lesson0073 import LESSON_0073
from lesson0074 import LESSON_0074
from lesson0075 import LESSON_0075
from lesson0076 import LESSON_0076
from lesson0077 import LESSON_0077
from lesson0078 import LESSON_0078
from lesson0079 import LESSON_0079
from lesson0080 import LESSON_0080
from lesson0081 import LESSON_0081
from lesson0082 import LESSON_0082
from lesson0083 import LESSON_0083
from lesson0084 import LESSON_0084
from lesson0085 import LESSON_0085
from lesson0086 import LESSON_0086
from lesson0087 import LESSON_0087
from lesson0088 import LESSON_0088
from lesson0089 import LESSON_0089
from lesson0090 import LESSON_0090
from lesson0091 import LESSON_0091
from lesson0092 import LESSON_0092
from lesson0093 import LESSON_0093
from lesson0094 import LESSON_0094
from lesson0095 import LESSON_0095
from lesson0096 import LESSON_0096
from lesson0097 import LESSON_0097
from lesson0098 import LESSON_0098
from lesson0099 import LESSON_0099
from lesson0100 import LESSON_0100
from lesson0101 import LESSON_0101
from lesson0102 import LESSON_0102
from lesson0103 import LESSON_0103
from lesson0104 import LESSON_0104
from lesson0105 import LESSON_0105
from lesson0106 import LESSON_0106
from lesson0107 import LESSON_0107
from lesson0108 import LESSON_0108

DELIVERY1 = [LESSON_0002, LESSON_0003, LESSON_0004, LESSON_0005, LESSON_0006, LESSON_0007]
DELIVERY2 = [LESSON_0008, LESSON_0009, LESSON_0010, LESSON_0011, LESSON_0012, LESSON_0013, LESSON_0014, LESSON_0015]
DELIVERY3 = [LESSON_0016, LESSON_0017, LESSON_0018, LESSON_0019, LESSON_0020, LESSON_0021, LESSON_0022, LESSON_0023]
DELIVERY4 = [LESSON_0024, LESSON_0025, LESSON_0026, LESSON_0027, LESSON_0028, LESSON_0029, LESSON_0030, LESSON_0031]
DELIVERY5 = [LESSON_0032, LESSON_0033, LESSON_0034, LESSON_0035, LESSON_0036, LESSON_0037, LESSON_0038, LESSON_0039]
DELIVERY6 = [LESSON_0040, LESSON_0041, LESSON_0042, LESSON_0043, LESSON_0044, LESSON_0045, LESSON_0046, LESSON_0047]
DELIVERY7 = [LESSON_0048, LESSON_0049, LESSON_0050, LESSON_0051, LESSON_0052, LESSON_0053, LESSON_0054, LESSON_0055]
DELIVERY8 = [LESSON_0056, LESSON_0057, LESSON_0058, LESSON_0059, LESSON_0060, LESSON_0061, LESSON_0062, LESSON_0063]
DELIVERY9 = [LESSON_0064, LESSON_0065, LESSON_0066, LESSON_0067, LESSON_0068, LESSON_0069, LESSON_0070, LESSON_0071, LESSON_0072]
DELIVERY10 = [LESSON_0073, LESSON_0074, LESSON_0075, LESSON_0076, LESSON_0077, LESSON_0078, LESSON_0079, LESSON_0080, LESSON_0081]
DELIVERY11 = [LESSON_0082, LESSON_0083, LESSON_0084, LESSON_0085, LESSON_0086, LESSON_0087, LESSON_0088, LESSON_0089, LESSON_0090]
DELIVERY12 = [LESSON_0091, LESSON_0092, LESSON_0093, LESSON_0094, LESSON_0095, LESSON_0096, LESSON_0097, LESSON_0098, LESSON_0099]

LESSONS_DELIVERY13 = [LESSON_0100, LESSON_0101, LESSON_0102, LESSON_0103, LESSON_0104, LESSON_0105, LESSON_0106, LESSON_0107, LESSON_0108]  # append more here as written; closes near ~37 pages

def main():
    doc = docx.Document()
    new_section_setup(doc)

    for idx, lesson in enumerate(LESSONS_DELIVERY13):
        if idx != 0:
            doc.add_page_break()
        add_title_block(doc)
        add_lesson_header_table(doc, lesson['cefr'], lesson['lesson_id'], lesson['en_title'], lesson['vi_title'])
        add_blank_spacer(doc)
        add_intro_paragraph(doc, lesson['intro_en'], lesson['intro_vi'])
        for speaker, en, vi in lesson['turns']:
            add_dialogue_line(doc, speaker, en, vi)

    lesson_range = f"{LESSONS_DELIVERY13[0]['lesson_id']}-{LESSONS_DELIVERY13[-1]['lesson_id']}"

    sec = doc.sections[0]
    ftr = sec.footer
    p = ftr.paragraphs[0]
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"EVERYDAY ENGLISH REFLEX  •  Lessons {lesson_range}  •  page ")
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    r.font.size = Pt(8)
    add_page_number_field(p)

    out_path = f"EVERYDAY_ENGLISH_REFLEX_LESSONS_{lesson_range}.docx"
    doc.save(out_path)
    print("Saved:", out_path)

    print("\n=== PER-LESSON QC (Delivery #13) ===")
    for lesson in LESSONS_DELIVERY13:
        r = qc_report(lesson)
        print(f"Lesson {r['lesson_id']}: English words={r['english_words']}, turns={r['turns']}, "
              f"speakers={sorted(r['speakers'])}, duplicate_lines={r['duplicate_lines']}")

    lesson0001 = load_lesson0001_as_dict("lesson0001_approved.docx")
    WHOLE_BOOK = [lesson0001] + DELIVERY1 + DELIVERY2 + DELIVERY3 + DELIVERY4 + DELIVERY5 + DELIVERY6 + DELIVERY7 + DELIVERY8 + DELIVERY9 + DELIVERY10 + DELIVERY11 + DELIVERY12 + LESSONS_DELIVERY13
    dups = cross_lesson_duplicate_check(WHOLE_BOOK)
    print(f"\nCross-lesson duplicate English lines (whole book, 0001 through {LESSONS_DELIVERY13[-1]['lesson_id']}):", dups)

    total_words = sum(lesson_word_count(l) for l in LESSONS_DELIVERY13)
    total_turns = sum(len(l['turns']) for l in LESSONS_DELIVERY13) + len(LESSONS_DELIVERY13)
    est_pages = total_turns / (133 / 5)
    print(f"\nTotal English learning words, Delivery #13 ({lesson_range}):", total_words)
    print(f"Structural page estimate for Delivery #13 so far: {est_pages:.1f} pages")

if __name__ == "__main__":
    main()
